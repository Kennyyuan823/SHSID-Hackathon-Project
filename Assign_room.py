"""Final room assignment: one fixed room per Section, sections share rooms
at different times, no room-time conflicts. Minimizes student walking cost.

Distance model:
  S1..S18 weighted ring, edges [1,1,3, 1*14, 3]
  S19-S20=1, S20-S21=1, S19-S21=2
  S1..S18 <-> S19..S21 = 10

Cost: a = 0 if d==0 else c + x*ln(d), summed over consecutive lessons.
Solved by simulated annealing with conflict penalty. c,x tuned on 80/20 split.

Outputs in ./output:
  rooms_layout.xlsx, rooms_by_slot.xlsx, metrics.xlsx,
  master_schedule.xlsx (with Room column)
"""
from __future__ import annotations
import math, random, sys
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUT = HERE / "output"
MASTER_CANDIDATES = [HERE/"master_schedule.xlsx", HERE/"output"/"master_schedule.xlsx",
                     Path.cwd()/"master_schedule.xlsx", Path.cwd()/"output"/"master_schedule.xlsx"]

ROOMS = [f"S{i}" for i in range(1, 22)]
RING_ORDER = [f"S{i}" for i in range(1, 19)]
RING_EDGE = [1, 1, 3] + [1]*14 + [3]

def _ring_dist(a, b):
    if a == b: return 0
    n = len(RING_ORDER)
    ia, ib = RING_ORDER.index(a), RING_ORDER.index(b)
    def fwd(x, y):
        t, k = 0, x
        while k != y:
            t += RING_EDGE[k % n]; k = (k+1) % n
        return t
    return min(fwd(ia, ib), fwd(ib, ia))

def build_distance():
    D = {}
    for a in ROOMS:
        for b in ROOMS:
            if a == b: D[(a,b)] = 0; continue
            ia, ib = int(a[1:]), int(b[1:])
            if ia <= 18 and ib <= 18:
                D[(a,b)] = _ring_dist(a,b)
            elif ia >= 19 and ib >= 19:
                D[(a,b)] = {(19,20):1,(20,19):1,(20,21):1,(21,20):1,(19,21):2,(21,19):2}.get((ia,ib),1)
            else:
                D[(a,b)] = 10
    return D

REQ = ["Student","Subject","Level","Section","Teacher","Weekly minutes","When"]
DAY_INDEX = {"Mon":0,"Tue":1,"Wed":2,"Thu":3,"Fri":4}
CLOCK_SLOTS = {"8:30-9:35":(0,1),"13:05-14:10":(5,6),"14:20-15:50":(7,8),
    "8:30-9:00":(0,),"9:05-9:35":(1,),"9:45-10:25":(2,),"10:35-11:15":(3,),
    "11:25-12:05":(4,),"13:05-13:35":(5,),"13:40-14:10":(6,),"14:20-15:00":(7,),"15:10-15:50":(8,)}

def find_master():
    for p in MASTER_CANDIDATES:
        if p.is_file(): return p
    raise SystemExit("找不到 master_schedule.xlsx（重命名后放脚本同目录）")

def header_map(ws):
    h = next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    return {str(n):i for i,n in enumerate(h) if n}

def parse_when(text):
    out=[]
    for piece in str(text or "").split(";"):
        piece=piece.strip()
        if not piece: continue
        week=None
        for tag in ("单周","双周"):
            if piece.endswith(tag): week=tag[0]; piece=piece[:-len(tag)].strip()
        day,clock=piece.split(None,1)
        for s in CLOCK_SLOTS[clock]: out.append((DAY_INDEX[day],s,week))
    return out

def load(master):
    print(f"[load] 读取 {master}", flush=True)
    wb=load_workbook(master,data_only=True,read_only=True,keep_links=False)
    cols=header_map(wb["Enrollments"])
    miss=[c for c in REQ if c not in cols]
    if miss: wb.close(); raise SystemExit(f"缺列 {miss}")
    sections={}; students=defaultdict(list)
    n=0
    for row in wb["Enrollments"].iter_rows(min_row=2,values_only=True):
        if not row or not row[cols["Student"]]: continue
        n+=1
        sid=str(row[cols["Student"]]); sec=str(row[cols["Section"]]); subj=row[cols["Subject"]]
        if sec not in sections:
            sections[sec]={"subject":subj,"level":row[cols["Level"]] or "",
                           "teacher":row[cols["Teacher"]],"when":row[cols["When"]] or "",
                           "block":row[cols["Block"]] if "Block" in cols and row[cols["Block"]] else "",
                           "meetings":parse_when(row[cols["When"]] or "")}
        for d,s,w in parse_when(row[cols["When"]] or ""):
            students[sid].append(((d,s,w or ""),sec))
    wb.close()
    for sid in students: students[sid].sort(key=lambda t:t[0])
    print(f"[load] 完成：{len(sections)} Section, {len(students)} 学生, {n} 行", flush=True)
    return sections, students

def section_slots(sections):
    return {sec: {(d,s,w or "") for d,s,w in info["meetings"]}
            for sec,info in sections.items()}

def pair_weights(students):
    sym=defaultdict(float)
    for _,lst in students.items():
        for (k1,s1),(k2,s2) in zip(lst,lst[1:]):
            if s1==s2: continue
            sym[(s1,s2)]+=1; sym[(s2,s1)]+=1
    return sym

def walk_cost(assign,w,D,c,x):
    t=0.0
    for (a,b),n in w.items():
        d=D[(assign[a],assign[b])]
        if d>0: t+=n*(c+x*math.log(d))
    return t

def conflicts(assign, sec_slots):
    by=defaultdict(int)
    for sec,room in assign.items():
        for t in sec_slots[sec]:
            by[(room,t)]+=1
    return sum(max(0,v-1) for v in by.values())

def initial_assign(sections, sec_slots, rng):
    names=sorted(sections,key=lambda s:-len(sec_slots[s]))
    used=defaultdict(set); assign={}
    for sec in names:
        times=sec_slots[sec]
        cand=[r for r in ROOMS if all((r,t) not in used for t in times)]
        r=rng.choice(cand) if cand else rng.choice(ROOMS)
        assign[sec]=r
        for t in times: used[(r,t)].add(sec)
    return assign

def anneal(sections,w,D,c,x,seed=0,iters=80000,penalty=1e5,tag=""):
    rng=random.Random(seed)
    names=list(sections)
    ss=section_slots(sections)
    assign=initial_assign(sections,ss,rng)
    cur=walk_cost(assign,w,D,c,x)+penalty*conflicts(assign,ss)
    best,ba=cur,dict(assign)
    T0,T1=5.0,0.01
    for it in range(iters):
        if it and it % 20000 == 0:
            print(f"  anneal{tag} it={it}/{iters} cur={cur:.1f}", flush=True)
        T=T0*(T1/T0)**(it/iters)
        a,b=rng.sample(names,2)
        ra,rb=assign[a],assign[b]
        if ra==rb: continue
        assign[a],assign[b]=rb,ra
        new=walk_cost(assign,w,D,c,x)+penalty*conflicts(assign,ss)
        if new<=cur or rng.random()<math.exp((cur-new)/max(T,1e-9)):
            cur=new
            if cur<best: best,ba=cur,dict(assign)
        else: assign[a],assign[b]=ra,rb
    return ba, walk_cost(ba,w,D,c,x), conflicts(ba,ss)

def peak_parallel(sections):
    by=defaultdict(set)
    for sec,info in sections.items():
        for d,s,w in info["meetings"]:
            by[(d,s,w or "")].add(sec)
    peak=max((len(v) for v in by.values()),default=0)
    print(f"[峰值] 同一时段并行 Section 最大数 = {peak}（教室 {len(ROOMS)} 个）", flush=True)
    if peak>len(ROOMS):
        print("[警告] 峰值 > 教室数，21 间不够，需要合班或加场地。", flush=True)
    return peak

def style_header(ws,n):
    f=PatternFill("solid",fgColor="1F4E79"); ft=Font(bold=True,color="FFFFFF")
    for c in range(1,n+1):
        cell=ws.cell(1,c); cell.fill=f; cell.font=ft
        cell.alignment=Alignment(horizontal="center",vertical="center")

def write_layout(D,assign,path):
    wb=Workbook(); ws=wb.active; ws.title="距离"
    ws.cell(1,1,"S\\T")
    for j,r in enumerate(ROOMS,2): ws.cell(1,j,r)
    for i,r in enumerate(ROOMS,2):
        ws.cell(i,1,r)
        for j,r2 in enumerate(ROOMS,2): ws.cell(i,j,D[(r,r2)])
    style_header(ws,len(ROOMS)+1)
    ws2=wb.create_sheet("分配"); ws2.append(["Section","Room"])
    for sec in sorted(assign): ws2.append([sec,assign[sec]])
    style_header(ws2,2); wb.save(path)

def write_by_slot(sections,assign,path):
    wb=Workbook(); ws=wb.active; ws.title="教室时段"
    header=["Day","Slot","Week"]+ROOMS; ws.append(header); style_header(ws,len(header))
    cellmap=defaultdict(lambda: defaultdict(list))
    for sec,info in sections.items():
        room=assign[sec]
        for d,s,w in info["meetings"]:
            cellmap[(d,s,w or "")][room].append(f"{sec} {info['subject']} {info['level']} / {info['teacher']}")
    dn=["Mon","Tue","Wed","Thu","Fri"]; row=2
    red=PatternFill("solid",fgColor="FFC7CE")
    for key in sorted(cellmap):
        d,s,w=key
        ws.cell(row,1,dn[d]); ws.cell(row,2,s); ws.cell(row,3,w or "")
        for j,r in enumerate(ROOMS,4):
            items=cellmap[key].get(r,[])
            ws.cell(row,j,"; ".join(items))
            if len(items)>1: ws.cell(row,j).fill=red
        row+=1
    for i in range(1,len(header)+1): ws.column_dimensions[get_column_letter(i)].width=14 if i<=3 else 28
    wb.save(path)

def write_metrics(students,assign,D,c,x,best,path,baseline,conf):
    wb=Workbook(); ws=wb.active; ws.title="总览"
    ws.append(["参数","值"])
    for k,v in [("c",c),("x",x),("优化后走班总代价 a",round(best,4)),
                ("随机基线 a",round(baseline,4)),
                ("改进%",round((baseline-best)/baseline*100,2) if baseline else 0),
                ("教室冲突数",conf),("学生数",len(students)),("Section 数",len(assign))]:
        ws.append([k,v])
    style_header(ws,2)
    ws2=wb.create_sheet("每生"); ws2.append(["Student","a_i","走班次数"])
    for sid,lst in sorted(students.items()):
        a=0.0; n=0
        for (k1,s1),(k2,s2) in zip(lst,lst[1:]):
            if s1==s2: continue
            d=D[(assign[s1],assign[s2])]
            if d>0: a+=c+x*math.log(d); n+=1
        ws2.append([sid,round(a,4),n])
    style_header(ws2,3); wb.save(path)

def fill_master_room(master,assign,path):
    wb=load_workbook(master); ws=wb["Enrollments"]; cols=header_map(ws)
    if "Room" not in cols:
        ws.cell(1,ws.max_column+1,"Room"); cols["Room"]=ws.max_column-1
    amap={str(k).strip():v for k,v in assign.items()}
    filled=0
    for row in range(2,ws.max_row+1):
        sec=ws.cell(row,cols["Section"]+1).value
        if sec is None: continue
        sec=str(sec).strip()
        if sec in amap:
            ws.cell(row,cols["Room"]+1,amap[sec]); filled+=1
    wb.save(path)
    print(f"[fill_master_room] 写入 {filled} 行 Room -> {path}", flush=True)

def main():
    print("[1] start", flush=True)
    OUT.mkdir(exist_ok=True)
    master=find_master()
    print(f"[2] master = {master}", flush=True)
    sections,students=load(master)
    peak_parallel(sections)
    weights=pair_weights(students)
    D=build_distance()

    ids=sorted(students); random.Random(1).shuffle(ids)
    cut=max(1,int(len(ids)*0.8)); train=set(ids[:cut]); val=set(ids[cut:]) or set(ids[:cut])
    w_train=pair_weights({s:students[s] for s in train})
    w_val  =pair_weights({s:students[s] for s in val})
    grid=[(c,x) for c in (0.5,1.0,2.0) for x in (0.5,1.0,2.0,3.0)]
    best_cfg=None
    for c,x in grid:
        assign_tr,_,conf=anneal(sections,w_train,D,c,x,seed=42,iters=30000,tag=f"({c},{x})")
        if conf>0:
            print(f"  (c={c}, x={x}) 冲突 {conf}，跳过", flush=True); continue
        v=walk_cost(assign_tr,w_val,D,c,x)
        print(f"  (c={c}, x={x}) val_a={v:.2f}", flush=True)
        if best_cfg is None or v<best_cfg[0]: best_cfg=(v,c,x)
    if best_cfg is None:
        print("[警告] 全部冲突，回退 c=1,x=1", flush=True); best_cfg=(0.0,1.0,1.0)
    _,c,x=best_cfg
    print(f"[选参] 最优 c={c}, x={x}", flush=True)

    assign,best,conf=anneal(sections,weights,D,c,x,seed=7,iters=60000,tag=" final")
    print(f"[最终退火完成] a={best:.2f}, conflicts={conf}", flush=True)

    ss=section_slots(sections)
    rand_assign=initial_assign(sections,ss,random.Random(7))
    baseline=walk_cost(rand_assign,weights,D,c,x)
    print(f"[基线] a={baseline:.2f}", flush=True)

    print("[写入] rooms_layout.xlsx ...", flush=True)
    write_layout(D,assign,OUT/"rooms_layout.xlsx")
    print("[写入] rooms_by_slot.xlsx ...", flush=True)
    write_by_slot(sections,assign,OUT/"rooms_by_slot.xlsx")
    print("[写入] metrics.xlsx ...", flush=True)
    write_metrics(students,assign,D,c,x,best,OUT/"metrics.xlsx",baseline,conf)
    print("[写入] master_schedule.xlsx (Room) ...", flush=True)
    fill_master_room(master,assign,OUT/"master_schedule.xlsx")
    print(f"c={c}, x={x}, a={best:.2f}, baseline={baseline:.2f}, "
          f"improve={(baseline-best)/baseline*100:.1f}%, conflicts={conf}", flush=True)
    print(OUT, flush=True)

if __name__=="__main__":
    main()