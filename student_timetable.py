"""Usage:
    python student_timetable.py S013
    python student_timetable.py S013 output/master_schedule.xlsx
"""
from __future__ import annotations
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
MASTER_NAME="master_schedule.xlsx"
DAYS=["周一","周二","周三","周四","周五"]
DAY_INDEX={"Mon":0,"Tue":1,"Wed":2,"Thu":3,"Fri":4}
CLOCK_SLOTS={"8:30-9:35":(0,1),"13:05-14:10":(5,6),"14:20-15:50":(7,8),
    "8:30-9:00":(0,),"9:05-9:35":(1,),"9:45-10:25":(2,),"10:35-11:15":(3,),
    "11:25-12:05":(4,),"13:05-13:35":(5,),"13:40-14:10":(6,),"14:20-15:00":(7,),"15:10-15:50":(8,)}
ROW_LABELS=["8:30-9:00","9:05-9:35","9:45-10:25","10:35-11:15","11:25-12:05",
            "12:05-13:05","13:05-13:35","13:40-14:10","14:20-15:00","15:10-15:50"]
ROW_SLOT=[0,1,2,3,4,None,5,6,7,8]
CN={"Math":"数学","Physics":"物理","Chemistry":"化学","Biology":"生物","English":"英语",
    "Chinese":"语文","History":"历史","CS":"计算机","Econ":"经济","Art":"美术","Music":"音乐",
    "PE":"体育","PBL":"PBL","CAS":"CAS","二外":"二外","Optional":"选修"}
REQ=["Student","Subject","Level","Section","Teacher","Weekly minutes","When"]

def argb(v): v=v.removeprefix("#"); return ("FF"+v if len(v)==6 else v).upper()

def find_master(explicit):
    if explicit:
        if not explicit.is_file(): raise SystemExit(f"找不到 {explicit}")
        return explicit
    for p in [HERE/"output"/MASTER_NAME,HERE/MASTER_NAME,
              Path.cwd()/"output"/MASTER_NAME,Path.cwd()/MASTER_NAME]:
        if p.is_file(): return p
    raise SystemExit("找不到 master_schedule.xlsx（先跑 assign_rooms.py）")

def normalize(s):
    t=s.strip().upper()
    if t.isdigit(): return f"S{int(t):03d}"
    if t.startswith("S") and t[1:].isdigit(): return f"S{int(t[1:]):03d}"
    return t

def header_map(ws):
    h=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    return {str(n):i for i,n in enumerate(h) if n}

def parse_when(text):
    out=[]
    for piece in str(text or "").split(";"):
        piece=piece.strip()
        if not piece: continue
        week=None
        for tag in ("单周","双周"):
            if piece.endswith(tag): week=tag[0]; piece=piece[:-len(tag)].strip()
        day,clock=piece.split(None,1)
        for s in CLOCK_SLOTS[clock]: out.append((DAY_INDEX[day],s,week))
    return out

def load(master,sid):
    wb=load_workbook(master,data_only=True,read_only=True)
    if "Enrollments" not in wb.sheetnames: wb.close(); raise SystemExit("缺 Enrollments")
    cols=header_map(wb["Enrollments"]); miss=[c for c in REQ if c not in cols]
    if miss: wb.close(); raise SystemExit(f"缺列 {miss}")
    has_room="Room" in cols
    if not has_room:
        print(f"[警告] {master} 没有 Room 列，先跑 assign_rooms.py。", file=sys.stderr)
    courses=[]; seen=set()
    for row in wb["Enrollments"].iter_rows(min_row=2,values_only=True):
        if not row or str(row[cols["Student"]])!=sid: continue
        sec=str(row[cols["Section"]])
        if sec in seen: continue
        seen.add(sec); subj=row[cols["Subject"]]
        courses.append({"subject":subj,"cn":CN.get(subj,subj),
            "level":row[cols["Level"]] or "","section":sec,"teacher":row[cols["Teacher"]],
            "room":(row[cols["Room"]] if has_room and row[cols["Room"]] else "(no room)"),
            "minutes":int(row[cols["Weekly minutes"]]),"when":row[cols["When"]] or "",
            "meetings":parse_when(row[cols["When"]] or "")})
    wb.close()
    if not courses: raise SystemExit(f"没有学生 {sid}")
    return courses

def lesson(c,week=None):
    pre={"单":"单周 ","双":"双周 "}.get(week,"")
    return f"{pre}{c['cn']} {c['level']}\n教室 {c['room']}\n老师 {c['teacher']}"

def write(path,sid,courses):
    thin=Border(*[Side(style="thin",color=argb("D0D5DD"))]*4)
    hf=PatternFill("solid",fgColor=argb("1F4E79")); hfont=Font(bold=True,color=argb("FFFFFF"),size=12)
    body=Font(size=11); small=Font(size=10); title=Font(bold=True,size=16,color=argb("1F4E79"))
    center=Alignment(horizontal="center",vertical="center",wrap_text=True)
    left=Alignment(horizontal="left",vertical="center",wrap_text=True)
    fills={"stem":PatternFill("solid",fgColor=argb("D6EAF8")),
           "double":PatternFill("solid",fgColor=argb("FDEBD0")),
           "optional":PatternFill("solid",fgColor=argb("F5EEF8")),
           "lunch":PatternFill("solid",fgColor=argb("E5E8E8")),
           "free":PatternFill("solid",fgColor=argb("F4F6F7")),
           "rot":PatternFill("solid",fgColor=argb("FCF3CF"))}
    grid=[[[] for _ in range(9)] for _ in range(5)]
    for c in courses:
        for d,s,w in c["meetings"]: grid[d][s].append((c,w))
    for d in range(5):
        for s in range(9):
            grid[d][s].sort(key=lambda t:(0 if t[0]["subject"]=="二外" else 1,t[1] or ""))
    wb=Workbook(); sh=wb.active; sh.title="课表"
    sh.merge_cells("A1:F1"); sh["A1"]=f"{sid} 的一周课表"; sh["A1"].font=title; sh["A1"].alignment=left
    sh.merge_cells("A2:F2"); sh["A2"]="每格：学科/level · 教室 · 老师；单/双周课按周次上。"
    sh["A2"].font=Font(size=11,color=argb("34495E")); sh["A2"].alignment=left
    sh.row_dimensions[1].height=24; sh.row_dimensions[2].height=20
    for col,h in enumerate(["时间",*DAYS],1):
        c=sh.cell(4,col,h); c.font=hfont; c.fill=hf; c.alignment=center; c.border=thin
    for r,label in enumerate(ROW_LABELS,5):
        slot=ROW_SLOT[r-5]
        tc=sh.cell(r,1,label+(" 午饭" if slot is None else ""))
        tc.font=body; tc.alignment=center; tc.border=thin
        if slot is None: tc.fill=fills["lunch"]
        for d in range(5):
            cell=sh.cell(r,d+2); cell.font=small; cell.alignment=center; cell.border=thin
            if slot is None: cell.value="午饭"; cell.fill=fills["lunch"]; continue
            if d==4 and slot in (7,8): cell.value="无课"; cell.fill=fills["free"]; continue
            items=grid[d][slot]
            if not items: cell.value=""; continue
            cell.value="\n".join(lesson(c,w) for c,w in items)
            subs={c["subject"] for c,_ in items}
            if subs<={"二外","CAS"}: cell.fill=fills["rot"]
            elif any(c["subject"] in {"Math","Physics","Chemistry","Biology"} for c,_ in items): cell.fill=fills["stem"]
            elif any(c["subject"]=="Optional" for c,_ in items): cell.fill=fills["optional"]
            elif slot in (0,1,5,6) or any(c["subject"]=="PBL" for c,_ in items): cell.fill=fills["double"]
        sh.row_dimensions[r].height=74
    for i,w in enumerate([16,24,24,24,24,24],1): sh.column_dimensions[get_column_letter(i)].width=w
    sh.freeze_panes="B5"; sh.page_setup.orientation="landscape"; sh.page_setup.fitToPage=True
    sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=1
    sh.sheet_properties.pageSetUpPr.fitToPage=True; sh.page_setup.paperSize=sh.PAPERSIZE_A4
    sh.print_title_rows="1:4"; sh.oddHeader.left.text=f"{sid} timetable"; sh.auto_filter.ref="A4:F14"
    lst=wb.create_sheet("课程")
    for c,h in enumerate(["学科","Level","教室","老师","周分钟","时间"],1):
        cell=lst.cell(1,c,h); cell.font=hfont; cell.fill=hf; cell.alignment=center; cell.border=thin
    ordered=sorted(courses,key=lambda c:(-c["minutes"],c["cn"],c["section"]))
    for r,c in enumerate(ordered,2):
        vals=[c["cn"],c["level"].strip(),c["room"],c["teacher"],c["minutes"],c["when"]]
        for col,v in enumerate(vals,1):
            cell=lst.cell(r,col,v); cell.font=body
            cell.alignment=left if col==6 else center; cell.border=thin
        lst.row_dimensions[r].height=30
    tr=len(ordered)+2
    lst.cell(tr,1,"合计").font=Font(bold=True)
    lst.cell(tr,5,sum(c["minutes"] for c in courses)).font=Font(bold=True)
    for col in range(1,7): lst.cell(tr,col).border=thin; lst.cell(tr,col).alignment=center
    for i,w in enumerate([14,18,14,14,12,78],1): lst.column_dimensions[get_column_letter(i)].width=w
    lst.auto_filter.ref=f"A1:F{len(ordered)+1}"; lst.freeze_panes="A2"
    path.parent.mkdir(parents=True,exist_ok=True); wb.save(path)

def main(argv=None):
    argv=list(sys.argv if argv is None else argv)
    sid=normalize(argv[1]) if len(argv)>1 and not argv[1].startswith("-") else normalize(input("Student id: "))
    master=find_master(Path(argv[2]) if len(argv)>2 and not argv[2].startswith("-") else None)
    courses=load(master,sid)
    dest=HERE/"output"/"students"/f"{sid}_timetable.xlsx"
    write(dest,sid,courses); print(dest)

if __name__=="__main__":
    main()